"""
magalu_loader.py
=================
Módulo para leitura da planilha de resultados trimestrais do Magazine Luiza
(RESULTADO_2T26_POR.xlsx) e organização dos dados de todas as guias em
estruturas (arrays / DataFrames) prontas para alimentar dashboards.

Uso básico:
-----------
    from magalu_loader import MagaluDataLoader

    loader = MagaluDataLoader("RESULTADO_2T26_POR.xlsx")

    loader.listar_guias()                       # nomes de todas as guias
    loader.listar_indicadores("1. Indicadores")  # linhas (indicadores) de uma guia
    loader.get_series("1. Indicadores", "EBITDA")# array (pandas Series) com o indicador
    loader.get_sheet_df("4. Balanço Patrimonial") # DataFrame completo da guia
    loader.buscar_indicador("margem")            # busca por indicador em todas as guias

Estrutura interna:
-------------------
Depois de carregado, todos os dados ficam disponíveis em:
    loader.dados[nome_da_guia]  -> pandas DataFrame (índice = indicador, colunas = período)
    loader.arrays[nome_da_guia][indicador] -> numpy array (apenas os valores)
    loader.colunas[nome_da_guia] -> lista de períodos (cabeçalho) daquela guia

Casos especiais tratados:
--------------------------
- A maioria das guias tem uma única linha de cabeçalho (períodos: '1T18', '2T18'...
  ou datas de fechamento de balanço) na linha 2 da planilha.
- Guia "15. DRE Proforma": cabeçalho em DUAS linhas (Segmento + Período). É
  carregada com colunas em MultiIndex (Segmento, Período).
- Guia "14. Luizacred - Carteira Atraso": cada período ocupa DUAS colunas
  (valor em R$ e % da carteira). É carregada com MultiIndex (Período, Tipo),
  onde Tipo é "Valor" ou "%".
- Indicadores com nomes repetidos dentro da mesma guia (ex.: "Balanço
  Patrimonial" tem "Títulos e Valores Mobiliários" no circulante e no não
  circulante) são desambiguados com o prefixo da seção em que aparecem,
  ex.: "ATIVO CIRCULANTE | Títulos e Valores Mobiliários".
"""

import re
from datetime import datetime
from collections import Counter

import numpy as np
import pandas as pd
import openpyxl

PADRAO_TRIMESTRE = re.compile(r'^\s*\dT\d{2}\s*$')


def _eh_periodo(valor):
    """Identifica se uma célula representa um rótulo de período (data ou '1T18')."""
    if isinstance(valor, datetime):
        return True
    if isinstance(valor, str) and PADRAO_TRIMESTRE.match(valor.strip()):
        return True
    return False


def _formatar_periodo(valor):
    """Converte datas em string 'dd/mm/aaaa' e mantém strings de trimestre como estão."""
    if isinstance(valor, datetime):
        return valor.strftime('%d/%m/%Y')
    if valor is None:
        return None
    return str(valor).strip()


class MagaluDataLoader:
    """Carrega todas as guias da planilha de resultados do Magazine Luiza."""

    def __init__(self, caminho_arquivo, carregar_ao_iniciar=True):
        self.caminho_arquivo = caminho_arquivo
        self.wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)

        self.dados = {}      # nome_guia -> DataFrame (índice=indicador, colunas=período)
        self.colunas = {}    # nome_guia -> lista de períodos/cabeçalho
        self.arrays = {}     # nome_guia -> {indicador: np.array}

        if carregar_ao_iniciar:
            self.carregar_tudo()

    # ------------------------------------------------------------------
    # Carregamento
    # ------------------------------------------------------------------
    def carregar_tudo(self):
        """Percorre todas as guias da planilha e monta as estruturas de dados."""
        for nome_guia in self.wb.sheetnames:
            ws = self.wb[nome_guia]

            if nome_guia.strip().startswith('15.'):
                df = self._carregar_guia_multi_segmento(ws)
            elif nome_guia.strip().startswith('14.'):
                df = self._carregar_guia_valor_percentual(ws)
            else:
                df = self._carregar_guia_padrao(ws)

            self.dados[nome_guia] = df
            self.colunas[nome_guia] = list(df.columns)
            self.arrays[nome_guia] = {
                indicador: df.loc[indicador].to_numpy()
                for indicador in df.index
            }
        return self.dados

    def _encontrar_linha_cabecalho(self, ws, max_linhas_busca=8):
        """Encontra a linha (1-indexada) onde estão os rótulos de período."""
        for i in range(1, max_linhas_busca + 1):
            valores = [c.value for c in ws[i]][1:]
            n_periodo = sum(1 for v in valores if _eh_periodo(v))
            if n_periodo >= 3:
                return i
        raise ValueError(f"Não foi possível localizar a linha de cabeçalho na guia '{ws.title}'.")

    def _carregar_guia_padrao(self, ws):
        """Guias com uma única linha de cabeçalho (a maioria das guias)."""
        linha_cab = self._encontrar_linha_cabecalho(ws)
        cabecalho_bruto = [c.value for c in ws[linha_cab]][1:]
        colunas = [_formatar_periodo(v) for v in cabecalho_bruto]
        n_col = len(colunas)

        indicadores = []
        valores = []
        secao_atual = None

        for linha in ws.iter_rows(min_row=linha_cab + 1, max_row=ws.max_row):
            rotulo = linha[0].value
            resto = [c.value for c in linha[1:1 + n_col]]

            if rotulo is None:
                continue  # linha em branco / separador

            rotulo_limpo = str(rotulo).strip()
            tem_dados = any(v is not None for v in resto)

            if not tem_dados:
                # linha "título de seção" (ex.: "ATIVO CIRCULANTE"), sem valores
                secao_atual = rotulo_limpo
                continue

            resto = resto + [None] * (n_col - len(resto))
            indicadores.append((secao_atual, rotulo_limpo))
            valores.append([np.nan if v is None else v for v in resto])

        rotulos_finais = self._desambiguar_rotulos(indicadores)
        df = pd.DataFrame(valores, index=rotulos_finais, columns=colunas)
        return df

    def _desambiguar_rotulos(self, indicadores):
        """Se um mesmo rótulo aparece mais de uma vez na guia, prefixa com a seção."""
        contagem = Counter(rot for _, rot in indicadores)
        finais = []
        for secao, rotulo in indicadores:
            if contagem[rotulo] > 1 and secao:
                finais.append(f"{secao} | {rotulo}")
            else:
                finais.append(rotulo)
        return finais

    def _carregar_guia_valor_percentual(self, ws):
        """Guia '14. Luizacred - Carteira Atraso': cada período = 2 colunas (Valor / %)."""
        linha_cab = self._encontrar_linha_cabecalho(ws)
        cabecalho_bruto = [c.value for c in ws[linha_cab]][1:]

        # preenche (forward-fill) os períodos que ficam "None" na 2ª coluna do par
        periodos = []
        ultimo = None
        for v in cabecalho_bruto:
            if _eh_periodo(v):
                ultimo = _formatar_periodo(v)
            periodos.append(ultimo)

        tipos = ['Valor' if i % 2 == 0 else '%' for i in range(len(periodos))]
        colunas = pd.MultiIndex.from_arrays([periodos, tipos], names=['Período', 'Tipo'])
        n_col = len(colunas)

        indicadores = []
        valores = []
        for linha in ws.iter_rows(min_row=linha_cab + 1, max_row=ws.max_row):
            rotulo = linha[0].value
            resto = [c.value for c in linha[1:1 + n_col]]
            if rotulo is None or not any(v is not None for v in resto):
                continue
            resto = resto + [None] * (n_col - len(resto))
            indicadores.append(str(rotulo).strip())
            valores.append([np.nan if v is None else v for v in resto])

        df = pd.DataFrame(valores, index=indicadores, columns=colunas)
        return df

    def _carregar_guia_multi_segmento(self, ws):
        """Guia '15. DRE Proforma': cabeçalho em 2 linhas (linha 2=Segmento, linha 3=Período)."""
        segmentos = [c.value for c in ws[2]][1:]
        periodos_brutos = [c.value for c in ws[3]][1:]
        periodos = [_formatar_periodo(v) for v in periodos_brutos]

        colunas = pd.MultiIndex.from_arrays(
            [[str(s).strip() if s else s for s in segmentos], periodos],
            names=['Segmento', 'Período']
        )
        n_col = len(colunas)

        indicadores = []
        valores = []
        for linha in ws.iter_rows(min_row=4, max_row=ws.max_row):
            rotulo = linha[0].value
            resto = [c.value for c in linha[1:1 + n_col]]
            if rotulo is None or not any(v is not None for v in resto):
                continue
            resto = resto + [None] * (n_col - len(resto))
            indicadores.append(str(rotulo).strip())
            valores.append([np.nan if v is None else v for v in resto])

        df = pd.DataFrame(valores, index=indicadores, columns=colunas)
        return df

    # ------------------------------------------------------------------
    # Consulta / acesso aos dados
    # ------------------------------------------------------------------
    def listar_guias(self):
        """Retorna a lista com os nomes de todas as guias carregadas."""
        return list(self.dados.keys())

    def listar_indicadores(self, guia):
        """Retorna a lista de indicadores (linhas) disponíveis em uma guia."""
        self._validar_guia(guia)
        return list(self.dados[guia].index)

    def get_sheet_df(self, guia):
        """Retorna o DataFrame completo de uma guia (índice=indicador, colunas=período)."""
        self._validar_guia(guia)
        return self.dados[guia]

    def get_series(self, guia, indicador):
        """Retorna uma pandas Series (array indexado pelo período) de um indicador."""
        self._validar_guia(guia)
        df = self.dados[guia]
        if indicador not in df.index:
            sugestoes = [i for i in df.index if indicador.lower() in str(i).lower()]
            msg = f"Indicador '{indicador}' não encontrado na guia '{guia}'."
            if sugestoes:
                msg += f" Você quis dizer: {sugestoes[:5]}?"
            raise KeyError(msg)
        return df.loc[indicador]

    def get_array(self, guia, indicador):
        """Retorna apenas o array numpy (sem rótulos de período) de um indicador."""
        return self.get_series(guia, indicador).to_numpy()

    def buscar_indicador(self, termo):
        """Procura um termo (case-insensitive) no nome dos indicadores de todas as guias.
        Retorna lista de tuplas (guia, indicador)."""
        termo = termo.lower()
        resultados = []
        for guia, df in self.dados.items():
            for indicador in df.index:
                if termo in str(indicador).lower():
                    resultados.append((guia, indicador))
        return resultados

    def _validar_guia(self, guia):
        if guia not in self.dados:
            sugestoes = [g for g in self.dados if guia.lower() in g.lower()]
            msg = f"Guia '{guia}' não encontrada."
            if sugestoes:
                msg += f" Você quis dizer: {sugestoes}?"
            else:
                msg += f" Guias disponíveis: {self.listar_guias()}"
            raise KeyError(msg)


# ----------------------------------------------------------------------
# Execução direta: pequeno teste/demonstração
# ----------------------------------------------------------------------
if __name__ == "__main__":
    loader = MagaluDataLoader("RESULTADO_2T26_POR.xlsx")

    print("Guias disponíveis:")
    for g in loader.listar_guias():
        print(" -", g)

    print("\nExemplo -> Indicadores da guia '1. Indicadores':")
    print(loader.listar_indicadores("1. Indicadores")[:10])

    print("\nExemplo -> Série de EBITDA:")
    print(loader.get_series("1. Indicadores", "EBITDA"))

    print("\nExemplo -> Buscar indicadores com 'margem':")
    print(loader.buscar_indicador("margem")[:10])
